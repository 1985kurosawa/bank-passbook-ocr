from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from typing import List, Dict, Any, Optional
from .config import LOW_ROW_CONFIDENCE

def _is_number(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool)

def generate_simple_to_file(rows: List[Dict[str, Any]], headers: List[str], outfile: str,
                            row_confidences: Optional[List[float]] = None,
                            low_threshold: float = LOW_ROW_CONFIDENCE):
    from pathlib import Path
    Path(outfile).parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "明細"

    header_fill = PatternFill(start_color="FFEEE8", end_color="FFEEE8", fill_type="solid")
    low_fill = PatternFill(start_color="FFF59D", end_color="FFF59D", fill_type="solid")

    ws.append(headers)
    for c in range(1, len(headers)+1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for i, r in enumerate(rows, start=2):
        ws.append([r.get(h, "") for h in headers])
        is_low = False
        if row_confidences and len(row_confidences) >= (i-1):
            is_low = (row_confidences[i-2] < low_threshold)
        if is_low:
            for c in range(1, len(headers)+1):
                ws.cell(row=i, column=c).fill = low_fill

        for c, h in enumerate(headers, start=1):
            v = r.get(h, "")
            if h in ("出金", "入金", "残高") and _is_number(v):
                ws.cell(row=i, column=c).number_format = "#,##0.00"
                ws.cell(row=i, column=c).alignment = Alignment(horizontal="right")

    ws.freeze_panes = "A2"

    for idx, h in enumerate(headers, start=1):
        max_len = max([len(str(h))] + [len(str(r.get(h, ""))) for r in rows]) if rows else len(h)
        ws.column_dimensions[get_column_letter(idx)].width = max(8, min(28, max_len + 2))

    wb.save(outfile)
    return outfile
